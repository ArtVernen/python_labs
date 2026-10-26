
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Преобразует CSV в XLSX с автоматической шириной столбцов."""
    csv_file = Path(csv_path)
    xlsx_file = Path(xlsx_path)

    if csv_file.suffix.lower() != ".csv" or xlsx_file.suffix.lower() != ".xlsx":
        raise ValueError("Неверное расширение файла")

    if not csv_file.is_file():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    with csv_file.open("r", newline="", encoding="utf-8") as file:
        rows = list(csv.reader(file))

    if len(rows) < 2 or not rows[0] or "" in rows[0]:
        raise ValueError("В CSV должны быть заголовки и данные")

    for row in rows:
        if len(row) != len(rows[0]):
            raise ValueError("Строки CSV разной длины")

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"

    for row in rows:
        sheet.append(row)

    for col in range(1, len(rows[0]) + 1):
        width = 8

        for row in rows:
            width = max(width, len(row[col - 1]))

        letter = get_column_letter(col)
        sheet.column_dimensions[letter].width = width

    xlsx_file.parent.mkdir(parents=True, exist_ok=True)
    book.save(xlsx_file)


if __name__ == "__main__":
    csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")
    print("CSV -> XLSX: готово")
